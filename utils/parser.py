"""
Module 1: Resume Upload and Parsing
Extracts raw text content from uploaded PDF or DOCX resumes.
"""

import os
import pdfplumber
import docx


def extract_text(filepath: str) -> str:
    ext = os.path.splitext(filepath)[1].lower()

    if ext == ".pdf":
        return _extract_from_pdf(filepath)
    elif ext == ".docx":
        return _extract_from_docx(filepath)
    else:
        raise ValueError("Unsupported file type. Please upload a PDF or DOCX resume.")


def _extract_from_pdf(filepath: str) -> str:
    text_chunks = []
    with pdfplumber.open(filepath) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text_chunks.append(page_text)
    return "\n".join(text_chunks)


def _extract_from_docx(filepath: str) -> str:
    document = docx.Document(filepath)
    text_chunks = [p.text for p in document.paragraphs if p.text.strip()]

    # also capture text inside tables (some resumes use table layouts)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text_chunks.append(cell.text)

    return "\n".join(text_chunks)
